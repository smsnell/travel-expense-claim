"""Validate acquisition assertions, retain original bytes, and parse source records.

This module performs no remote access. Hashes establish byte integrity only.
"""
import hashlib
import json
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path, PurePosixPath
import re
import shutil
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PyPdfError


class SourceError(ValueError):
    """Source evidence is unavailable, malformed, incomplete, or incompatible."""


URLS = {
    'initial': 'https://drive.google.com/file/d/13sRj6s_kMeP1hsYGstWWmfkb6SCN1YOS/view',
    'updates': 'https://drive.google.com/file/d/16829RwrGrMCK9oq5TnWuytFewGk_8xQP/view',
    'receipts': 'https://drive.google.com/file/d/1Ynn1VLH-faRIy91xy95rOWwCBhy5mjjM/view',
    'registers': 'https://docs.google.com/spreadsheets/d/1338nvKD6rgAdLgGRhYCMMlQyDNLK0sE8pvFWrZAaaxg/edit',
    'policy-native': 'https://private-pecorino-70e.notion.site/Project-D-Alderbridge-travel-reimbursement-policy-POL-2026-1-3d80b700541e81ed93ebd386abc1dcb0',
}
TABS = ('People', 'Trips', 'Merchant payments', 'FX', 'Caps', 'Budget',
        'Review ledger', 'Review packets', 'Finance activity', 'Travel cancellations')
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def _require(condition, message):
    if not condition:
        raise SourceError(message)


def _hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _inside(root, relative):
    _require(isinstance(relative, str) and relative and not Path(relative).is_absolute(),
             'Source path must be nonempty and relative')
    path = (root / relative).resolve()
    _require(path.is_relative_to(root.resolve()), f'Source path escapes directory: {relative}')
    _require(path.is_file(), f'Source file missing: {relative}')
    return path


def _utc(value):
    _require(isinstance(value, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|\+00:00)', value), 'observed_at must be an ISO UTC timestamp')
    try:
        return datetime.fromisoformat(value.replace('Z', '+00:00'))
    except ValueError as exc:
        raise SourceError(f'Invalid observed_at: {value}') from exc


def _field(text, label):
    match = re.search(r'^' + re.escape(label) + r':\s*([^\n]*)', text, re.M)
    _require(match and match[1].strip(), f'Missing card field: {label}')
    return match[1].strip()


def _day(value):
    try:
        _require(date.fromisoformat(value).isoformat() == value, f'Invalid date: {value}')
    except ValueError as exc:
        raise SourceError(f'Invalid date: {value}') from exc
    return value


def _money(value):
    _require(bool(re.fullmatch(r'-?\d+(?:\.\d+)?', value)), f'Invalid decimal: {value}')
    _require(Decimal(value).is_finite(), f'Invalid decimal: {value}')
    return value


def _list(value):
    return [] if value == 'None' else [s.strip() for s in re.split(r'[,;\n]', value) if s.strip()]


def _claim(text, source_id, locator):
    result = {'_source_id': source_id, '_locator': locator}
    fields = {'claim_id': 'Claim reference', 'revision': 'Revision', 'batch_id': 'Arrival batch',
              'employee_id': 'Employee', 'trip_id': 'Trip reference', 'trip_revision': 'Trip revision',
              'submitted': 'Submitted date', 'status': 'Status', 'reason': 'Claim reason',
              'related_claim_ids': 'Related earlier claims'}
    result.update({key: _field(text, label) for key, label in fields.items()})
    for key in ('revision', 'trip_revision'):
        _require(result[key].isdigit() and int(result[key]) > 0, f'Invalid {key}')
        result[key] = int(result[key])
    result['reason'] = text.split('Claim reason:', 1)[1].split('Related earlier claims:', 1)[0].strip().replace('\n', ' ')
    _day(result['submitted'])
    result['related_claim_ids'] = _list(result['related_claim_ids'])
    _require('Expense details' in text and 'Finance exception' in text, 'Incomplete claim card')
    line_text = re.split(r'Original amount\s+Currency', text, maxsplit=1)
    _require(len(line_text) == 2, 'Missing expense table header')
    tokens = line_text[1].split('Finance exception', 1)[0].split()
    _require(len(tokens) % 6 == 0, 'Incomplete expense table row')
    lines = [' '.join(tokens[i:i+6]) for i in range(0, len(tokens), 6)]
    result['lines'] = []
    for number, line in enumerate(lines, 1):
        match = re.fullmatch(r'(\d+)\s+(\S+)\s+(\S+)\s+(\d{4}-\d{2}-\d{2})\s+(-?\d+(?:\.\d+)?)\s+([A-Z]{3})', line.strip())
        _require(match and int(match[1]) == number, f'Invalid expense line: {line}')
        result['lines'].append(dict(zip(('cost_id', 'receipt_id', 'first_submitted', 'original_amount', 'currency'), match.groups()[1:])))
        _day(match[4]); _money(match[5])
    _require(result['lines'], 'Claim has no expense lines')
    if 'Finance exception: Not supplied' not in text:
        labels = {'actor': 'Authorizing Finance officer', 'claim_id': 'Applies to claim',
                  'revision': 'Applies to revision', 'cost_id': 'Expense reference',
                  'allowed_original_amount': 'Allowed original amount', 'covered_issues': 'Covered issues'}
        exception_text = text.split('Finance exception', 1)[1]
        exception = {key: _field(exception_text, label) for key, label in labels.items()}
        _require(exception['revision'].isdigit(), 'Invalid exception revision')
        exception['revision'] = int(exception['revision'])
        exception['covered_issues'] = _list(exception['covered_issues'])
        exception['reason'] = exception_text.split('Exception reason:', 1)[-1].strip().replace('\n', ' ')
        _require('Exception reason:' in exception_text and exception['reason'], 'Missing exception reason')
        _money(exception['allowed_original_amount'])
        if 'Travel cancellation reference:' in exception_text:
            exception['cancellation_id'] = _field(exception_text, 'Travel cancellation reference')
        result['exception'] = exception
    return result


def _receipt(text, source_id, locator):
    labels = {'receipt_id': 'Receipt reference', 'cost_id': 'Expense reference', 'employee_id': 'Employee',
              'trip_id': 'Trip reference', 'category': 'Category', 'currency': 'Currency',
              'gross': 'Gross amount', 'units': 'Documented units', 'issued': 'Issued date', 'description': 'Description'}
    record = {key: _field(text, label) for key, label in labels.items()}
    record['description'] = text.split('Description:', 1)[1].split('Issued date:', 1)[0].strip().replace('\n', ' ')
    _require(record['units'].isdigit() and int(record['units']) > 0, 'Invalid receipt units')
    record['units'] = int(record['units'])
    _money(record['gross']); _day(record['issued'])
    record.update(_source_id=source_id, _locator=locator)
    return record


def _pdf(path, source_id, version):
    reader = PdfReader(path)
    records = []
    for number, page in enumerate(reader.pages, 1):
        text = page.extract_text().replace('\r', '')
        marker = re.findall(r'(SRC-[\w-]+)\s+Page (\d+) of (\d+)', text)
        _require(marker == [(version, str(number), str(len(reader.pages)))],
                 f'{source_id}: inconsistent version/page coverage at page {number}')
        label = 'Receipt reference:' if source_id == 'receipts' else 'Claim reference:'
        parts = re.split(r'(?m)(?=^' + re.escape(label) + ')', text)
        _require(len(parts) > 1, f'{source_id}: no cards on page {number}')
        for index, part in enumerate(parts[1:], 1):
            parser = _receipt if source_id == 'receipts' else _claim
            records.append(parser(part, source_id, f'page {number}, card {index}'))
    _require(records, f'{source_id}: empty PDF')
    identities = [(r.get('receipt_id') or r.get('claim_id'), r.get('revision')) for r in records]
    _require(len(set(identities)) == len(identities), f'{source_id}: duplicate cards')
    return records, {'page_count': len(reader.pages), 'record_count': len(records), 'page_markers': 'all physical pages matched'}


def _workbook(path, version):
    workbook = load_workbook(path, data_only=False)
    _require(tuple(workbook.sheetnames) == TABS, 'Workbook must contain all ten native tabs in order')
    tables, inventory = {}, {}
    with zipfile.ZipFile(path) as archive:
        wb = ET.fromstring(archive.read('xl/workbook.xml'))
        rels = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {r.get('Id'): r.get('Target') for r in rels}
        sheets = wb.find('m:sheets', NS)
        for sheet, node in zip(workbook.worksheets, sheets):
            target = targets[node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
            target = target.lstrip('/') if target.startswith('/') else str(PurePosixPath('xl') / target)
            xml = ET.fromstring(archive.read(target))
            raw = {c.get('r'): c.find('m:v', NS).text for c in xml.findall('.//m:c', NS) if c.find('m:v', NS) is not None}
            _require(version in str(sheet['A2'].value), f'{sheet.title}: source version conflict')
            headers = [c.value for c in sheet[5]]
            while headers and headers[-1] is None:
                headers.pop()
            _require(headers and all(isinstance(h, str) and h for h in headers) and len(headers) == len(set(headers)), f'{sheet.title}: invalid headers')
            records = []
            for row in sheet.iter_rows(min_row=6):
                if all(c.value is None for c in row):
                    continue
                _require(all(c.value is None for c in row[len(headers):]), f'{sheet.title}: values outside headers')
                record = {'_source_id': 'registers', '_locator': f'{sheet.title}!A{row[0].row}:{row[len(headers)-1].column_letter}{row[0].row}'}
                for header, cell in zip(headers, row):
                    value = cell.value
                    _require(cell.data_type != 'f', f'{sheet.title}!{cell.coordinate}: formulas are unsupported')
                    if value is None:
                        value = ''
                    elif isinstance(value, datetime):
                        value = value.replace(tzinfo=timezone.utc).isoformat().replace('+00:00', 'Z') if '(UTC)' in header else value.date().isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    elif isinstance(value, bool):
                        value = str(value)
                    elif isinstance(value, (int, float)):
                        value = raw[cell.coordinate]
                        if header in ('Arrival batch', 'Revision', 'Trip revision', 'Permit revision', 'Documented units') and Decimal(value) == Decimal(value).to_integral_value():
                            value = str(int(Decimal(value)))
                    else:
                        value = str(value)
                    record[header] = value
                records.append(record)
            _require(records, f'{sheet.title}: empty native table')
            tables[sheet.title] = records
            inventory[sheet.title] = {'headers': headers, 'records': len(records), 'last_row': records[-1]['_locator']}
    workbook.close()
    return tables, inventory


def _policy(path):
    text = path.read_text(encoding='utf-8')
    versions = set(re.findall(r'POL-\d{4}\.\d+', text))
    _require(versions == {'POL-2026.2'}, 'Policy version missing or conflicting; Operations owner must resolve')
    for term in ('Case clock:', 'receipt', 'merchant', 'cap', 'Finance', 'covered_issues', 'cancellation', 'reimbursement', 'closure', 'administration', 'supervisor', 'budget', 'director', 'refund', 'request', 'accepted', 'dispatch', 'reconciliation', 'revision', 'settled'):
        _require(term.lower() in text.lower(), f'Policy incomplete: missing {term}; Operations owner must resolve')
    _require(len(text) > 4000 and text.rstrip().endswith(('.', '`')), 'Policy appears truncated; Operations owner must resolve')
    return 'POL-2026.2'


def prepare_sources(acquisition_path, run_dir, policy_file, sequencing_file):
    """Validate a fresh acquisition handoff and bind exact copied inputs to a run.

    Failed validation raises SourceError. No downloaded fixture is treated as fresh.
    """
    acquisition_path, run_dir = Path(acquisition_path), Path(run_dir)
    _require(run_dir.is_dir() and not any(run_dir.iterdir()), 'run_dir must already exist and be empty')
    try:
        handoff = json.loads(acquisition_path.read_text(encoding='utf-8'))
        _require(handoff.get('mode') == 'fresh', 'G1 accepts only explicitly fresh acquisition handoffs')
        for key in ('acquisition_id', 'case_clock', 'source_version'):
            _require(handoff.get(key), f'Missing acquisition {key}')
        entries = handoff.get('files', [])
        _require(len(entries) == 5 and {e.get('source_id') for e in entries} == set(URLS), 'Need exactly five native source observations')
        prepared = []
        for entry in entries:
            sid = entry['source_id']
            for key in ('path', 'sha256', 'status', 'error', 'url', 'locator', 'observed_at', 'version', 'acquisition_method', 'completeness'):
                _require(key in entry, f'{sid}: missing {key}')
            _require(entry['url'] == URLS[sid], f'{sid}: unexpected native URL')
            _utc(entry['observed_at'])
            _require(entry['locator'] and entry['acquisition_method'] and entry['completeness'], f'{sid}: missing acquisition evidence')
            result = dict(entry)
            if entry['status'] == 'unavailable':
                _require(entry['path'] is None and entry['sha256'] is None and entry['error'], f'{sid}: invalid unavailable observation')
                _require(sid == 'policy-native', f'{sid}: required records unavailable; Operations owner must restore access')
            else:
                _require(entry['status'] == 'read' and entry['error'] is None, f'{sid}: invalid read status')
                path = _inside(acquisition_path.parent, entry['path'])
                _require(_hash(path) == entry['sha256'], f'{sid}: SHA-256 mismatch')
                _require(entry['version'], f'{sid}: missing observed source version')
                if sid == 'registers':
                    _require(entry['version'] == handoff['source_version'], f'{sid}: source version conflict')
                    _, result['verified_coverage'] = _workbook(path, handoff['source_version'])
                elif sid == 'policy-native':
                    _require(entry['version'] == _policy(path), 'Native policy version conflict')
                else:
                    _require(entry['version'] == handoff['source_version'], f'{sid}: source version conflict')
                    _, result['verified_coverage'] = _pdf(path, sid, handoff['source_version'])
                result['_input_path'] = path
            prepared.append(result)
        native = next(e for e in prepared if e['source_id'] == 'policy-native')
        _require(policy_file or native['status'] == 'read', 'Explicit policy_file required when native policy unavailable')
        if policy_file:
            policy_path = Path(policy_file)
            version = _policy(policy_path)
            if native['status'] == 'read':
                _require(policy_path.read_bytes() == native['_input_path'].read_bytes(), 'Native/manual policy differ; Operations owner must resolve')
            prepared.append(dict(source_id='policy-manual', _input_path=policy_path, path=None, sha256=_hash(policy_path), status='read', error=None, url=URLS['policy-native'], locator='Complete owner-supplied manual policy', observed_at=datetime.now(timezone.utc).isoformat(), version=version, acquisition_method='manual_file', origin=str(policy_path), export_time=None, live_equivalence='unverified'))
        sequence_path = Path(sequencing_file)
        _require(sequence_path.is_file() and sequence_path.stat().st_size > 0, 'Supplemental sequencing evidence required')
        prepared.append(dict(source_id='sequencing', _input_path=sequence_path, path=None, sha256=_hash(sequence_path), status='read', error=None, url=None, locator='Owner-confirmed interview sequencing evidence', observed_at=datetime.now(timezone.utc).isoformat(), version=None, acquisition_method='supplemental_local_evidence', origin=str(sequence_path)))
        # All evidence is validated before any run files are written.
        (run_dir / 'inputs').mkdir()
        for entry in prepared:
            path = entry.pop('_input_path', None)
            if path:
                destination = run_dir / 'inputs' / (entry['source_id'] + path.suffix)
                shutil.copyfile(path, destination)
                _require(_hash(destination) == entry['sha256'], 'Input changed during retention')
                entry['path'] = destination.relative_to(run_dir).as_posix()
        manifest = {key: handoff[key] for key in ('acquisition_id', 'mode', 'case_clock', 'source_version')}
        manifest.update(run_id=run_dir.name, parser_version='source-io/1', files=prepared,
                        provenance_note='Acquisition observations are operator/tool assertions; byte hashes do not authenticate remote origin.')
        (run_dir / 'sources.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
        return manifest
    except SourceError:
        raise
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile, ET.ParseError, PyPdfError, InvalidFileException) as exc:
        raise SourceError(f'Invalid source acquisition: {exc}') from exc


def normalize(run_dir):
    """Parse only the retained, hash-verified sources; return full source records."""
    run_dir = Path(run_dir)
    try:
        manifest = json.loads((run_dir / 'sources.json').read_text(encoding='utf-8'))
        result = {'claims': [], 'receipts': [], 'tables': {}}
        for entry in manifest['files']:
            if entry['status'] != 'read':
                continue
            path = _inside(run_dir, entry['path'])
            _require(_hash(path) == entry['sha256'], f"{entry['source_id']}: retained input SHA-256 mismatch")
            sid = entry['source_id']
            if sid in ('initial', 'updates', 'receipts'):
                records, _ = _pdf(path, sid, manifest['source_version'])
                result['receipts' if sid == 'receipts' else 'claims'].extend(records)
            elif sid == 'registers':
                result['tables'], _ = _workbook(path, manifest['source_version'])
        return result
    except SourceError:
        raise
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile, ET.ParseError, PyPdfError, InvalidFileException) as exc:
        raise SourceError(f'Cannot normalize retained sources: {exc}') from exc
